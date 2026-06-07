# main.py

import os
import sys

_PROJECT_ROOT = os.path.dirname(os.path.abspath(__file__))
if _PROJECT_ROOT not in sys.path:
    sys.path.insert(0, _PROJECT_ROOT)

import signal
import traceback
from typing import Dict, List, Optional
from sector_ui import SectorUI
try:
    import pandas as pd
    from rich.console  import Console
    from rich.panel    import Panel
    from rich.prompt   import Confirm, IntPrompt, Prompt
    from rich.progress import Progress, SpinnerColumn, BarColumn, TextColumn, TaskProgressColumn
    from rich.text     import Text
    from rich.table    import Table
    from rich          import box
except ImportError as e:
    print(f"\n❌ Missing package: {e}")
    print("   Run: pip install -r requirements.txt\n")
    sys.exit(1)

try:
    from config              import ScraperConfig
    from logger_setup        import get_logger
    from scraper             import IndianStockScraper
    from data_manager        import DataManager
    from filter_ui           import FilterUI, print_dataframe, print_statistics
    from filters             import FilterPipeline
    from company_selector    import CompanySelector
    from company_data        import ALL_COMPANIES, symbol_to_info
    from financial_scraper   import ScreenerFinancialScraper
except ImportError as e:
    print(f"\n❌ Local import error: {e}")
    traceback.print_exc()
    sys.exit(1)

console = Console()


def graceful_exit(sig, frame):
    console.print("\n[yellow]⚠  Interrupted — Goodbye![/yellow]")
    sys.exit(0)

signal.signal(signal.SIGINT,  graceful_exit)
signal.signal(signal.SIGTERM, graceful_exit)


# ─────────────────────────────────────────────────────────────────────────────
# Banner
# ─────────────────────────────────────────────────────────────────────────────

def show_banner():
    console.print(Panel(
        Text.assemble(
            ("  🇮🇳  Indian Stock Market Scraper  📈\n\n",         "bold white"),
            ("  🏢 Nifty 50 · Midcap · Smallcap (147 companies)\n","cyan"),
            ("  💰 Revenue · PAT · EBITDA (Last 3 Years)\n",        "green"),
            ("  🔍 16 Filter Types · Presets · Excel Export\n",     "yellow"),
            ("  📊 Sources: NSE · Yahoo · Screener.in",             "dim"),
        ),
        border_style="bright_blue",
        padding=(1, 6),
    ))


# ─────────────────────────────────────────────────────────────────────────────
# Scrape with company selection
# ─────────────────────────────────────────────────────────────────────────────

def run_scrape_with_selection(config: ScraperConfig,
                               logger) -> tuple:
    """
    Step 1: Let user pick companies.
    Step 2: Scrape live price data.
    Step 3: Optionally scrape financials (PAT, Revenue).
    Returns (price_df, financials_df)
    """
    dm = DataManager(config)

    # ── Step 1: Company Selection ─────────────────────────────────────────────
    console.print(Panel(
        "[bold cyan]Step 1: Select Companies to Scrape[/bold cyan]",
        border_style="cyan",
    ))

    selector = CompanySelector()
    selected = selector.run()   # {symbol: {name, sector, index}}

    if not selected:
        console.print("[red]No companies selected. Aborting.[/red]")
        return None, None

    # Update config watchlist with selected symbols
    config.WATCHLIST = list(selected.keys())

    # Update sector map
    for sym, info in selected.items():
        config.SECTOR_MAP[sym] = info.get("sector", "Unknown")

    console.print(f"\n  [green]✔ {len(selected)} companies selected[/green]")

    # ── Step 2: Choose what to scrape ─────────────────────────────────────────
    console.print(Panel(
        "[bold cyan]Step 2: What data to fetch?[/bold cyan]",
        border_style="cyan",
    ))

    console.print("  [1] Live prices only  (fast, ~1 min)")
    console.print("  [2] Live prices + Financial data  (Revenue, PAT, EBITDA)")
    console.print("  [3] Financial data only  (Revenue, PAT, EBITDA)")
    console.print("  [Q] Cancel\n")

    data_choice = Prompt.ask("  Choice", choices=["1","2","3","Q"],
                              default="1").upper()
    if data_choice == "Q":
        return None, None

    price_df      = None
    financials_df = None

    # ── Step 3: Scrape live prices ────────────────────────────────────────────
    if data_choice in ("1", "2"):
        console.print(Panel(
            "[bold green]Scraping live price data…[/bold green]",
            border_style="green",
        ))

        with Progress(
            SpinnerColumn(),
            TextColumn("[bold blue]{task.description}"),
            BarColumn(),
            TaskProgressColumn(),
            console=console,
        ) as progress:
            task = progress.add_task("Fetching prices…", total=None)
            scraper = IndianStockScraper(config)
            records = scraper.run_all()
            progress.update(task, completed=True)

        if records:
            price_df = dm.to_dataframe(records)
            console.print(
                f"  [green]✔ {len(price_df)} price records fetched[/green]"
            )
        else:
            console.print("  [red]⚠  No price data returned[/red]")

    # ── Step 4: Scrape financials ─────────────────────────────────────────────
    if data_choice in ("2", "3"):
        console.print(Panel(
            "[bold green]Scraping financial data (Screener.in)…[/bold green]\n"
            "[dim]  This may take a few minutes — respecting rate limits[/dim]",
            border_style="green",
        ))

        fin_scraper = ScreenerFinancialScraper(config, logger)
        fin_records = []
        symbols     = list(selected.keys())
        total       = len(symbols)

        with Progress(
            SpinnerColumn(),
            TextColumn("[bold blue]{task.description}"),
            BarColumn(),
            TaskProgressColumn(),
            console=console,
        ) as progress:
            task = progress.add_task("Fetching financials…", total=total)

            for i, sym in enumerate(symbols, 1):
                progress.update(
                    task,
                    description=f"[bold blue]{sym} ({i}/{total})…",
                    completed=i,
                )
                data = fin_scraper.scrape_company(sym)
                fin_records.append(data)

        if fin_records:
            financials_df = pd.DataFrame(fin_records)
            console.print(
                f"  [green]✔ {len(financials_df)} financial records fetched[/green]"
            )

    return price_df, financials_df


# ─────────────────────────────────────────────────────────────────────────────
# Save to Excel with multiple sheets
# ─────────────────────────────────────────────────────────────────────────────

def save_to_excel(config: ScraperConfig, logger,
                  price_df: Optional[pd.DataFrame],
                  financials_df: Optional[pd.DataFrame],
                  filtered_df: Optional[pd.DataFrame] = None):
    """Save all data to Excel with dedicated sheets."""

    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils  import get_column_letter
    from openpyxl.chart  import BarChart, Reference

    path = config.excel_path

    try:
        # Load or create workbook
        append = config.APPEND_MODE and os.path.exists(path)
        wb     = openpyxl.load_workbook(path) if append else openpyxl.Workbook()

        if not append and "Sheet" in wb.sheetnames:
            del wb["Sheet"]

        thin   = Side(style="thin", color="CCCCCC")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        def style_header(ws, num_cols, bg="1F4E79"):
            for ci in range(1, num_cols + 1):
                c = ws.cell(row=1, column=ci)
                c.font      = Font(bold=True, color="FFFFFF", size=11)
                c.fill      = PatternFill("solid", fgColor=bg)
                c.alignment = Alignment(horizontal="center", wrap_text=True)
                c.border    = border
            ws.row_dimensions[1].height = 22

        def write_df(ws, df, cat_colour=True):
            cat_colours = {
                "Top Gainer": "C6EFCE", "Top Loser":  "FFC7CE",
                "Index"     : "DDEBF7", "Watchlist":  "FFEB9C",
            }
            # Header
            for ci, col in enumerate(df.columns, 1):
                ws.cell(row=1, column=ci, value=col)
            style_header(ws, len(df.columns))

            # Data
            for ri, (_, row) in enumerate(df.iterrows(), 2):
                cat  = str(row.get("Category", ""))
                fill = PatternFill("solid",
                                   fgColor=cat_colours.get(cat, "FFFFFF")) \
                       if cat_colour else None
                for ci, val in enumerate(row, 1):
                    c           = ws.cell(row=ri, column=ci, value=str(val))
                    c.border    = border
                    c.alignment = Alignment(horizontal="left")
                    if fill:
                        c.fill = fill

            # Auto-width
            for col in ws.columns:
                w = max((len(str(c.value or "")) for c in col), default=10)
                ws.column_dimensions[
                    get_column_letter(col[0].column)
                ].width = min(w + 3, 40)

            ws.freeze_panes   = "A2"
            ws.auto_filter.ref = ws.dimensions

        # ── Sheet: Live Prices ────────────────────────────────────────────────
        if price_df is not None and not price_df.empty:
            for sheet_name, cat in [
                ("📋 All Prices",  None),
                ("📈 Indices",     "Index"),
                ("🚀 Gainers",     "Top Gainer"),
                ("📉 Losers",      "Top Loser"),
                ("👁 Watchlist",   "Watchlist"),
            ]:
                df_sub = price_df if cat is None else price_df[price_df["Category"] == cat]
                if df_sub.empty:
                    continue
                if sheet_name in wb.sheetnames:
                    del wb[sheet_name]
                ws = wb.create_sheet(sheet_name)
                write_df(ws, df_sub)
                logger.info(f"  Sheet '{sheet_name}': {len(df_sub)} rows")

        # ── Sheet: Financials ─────────────────────────────────────────────────
        if financials_df is not None and not financials_df.empty:
            FIN_SHEET = "💰 Financials"
            if FIN_SHEET in wb.sheetnames:
                del wb[FIN_SHEET]
            ws_fin = wb.create_sheet(FIN_SHEET)
            write_df(ws_fin, financials_df, cat_colour=False)
            style_header(ws_fin, len(financials_df.columns), bg="1F6B3B")

            # ── Revenue Chart ─────────────────────────────────────────────────
            try:
                chart = BarChart()
                chart.type    = "col"
                chart.title   = "Revenue Last 3 Years (₹ Cr)"
                chart.y_axis.title = "Revenue (₹ Cr)"
                chart.x_axis.title = "Company"

                # Find Revenue Y1/Y2/Y3 columns
                cols      = list(financials_df.columns)
                rev_cols  = [i+1 for i, c in enumerate(cols) if "Revenue Y" in c and "Period" not in c]
                sym_col   = cols.index("Symbol") + 1

                if rev_cols:
                    data = Reference(
                        ws_fin,
                        min_col=min(rev_cols),
                        max_col=max(rev_cols),
                        min_row=1,
                        max_row=min(len(financials_df) + 1, 20),
                    )
                    cats = Reference(
                        ws_fin,
                        min_col=sym_col,
                        min_row=2,
                        max_row=min(len(financials_df) + 1, 20),
                    )
                    chart.add_data(data, titles_from_data=True)
                    chart.set_categories(cats)
                    chart.shape = 4
                    chart.width = 30
                    chart.height = 15
                    ws_fin.add_chart(chart, "T2")

            except Exception as ce:
                logger.warning(f"  Chart creation failed: {ce}")

        # ── Sheet: Filtered ───────────────────────────────────────────────────
        if filtered_df is not None and not filtered_df.empty:
            FILT_SHEET = "🔍 Filtered"
            if FILT_SHEET in wb.sheetnames:
                del wb[FILT_SHEET]
            ws_f = wb.create_sheet(FILT_SHEET)
            write_df(ws_f, filtered_df, cat_colour=False)

        # ── Sheet: Summary ────────────────────────────────────────────────────
        SUM_SHEET = "📊 Summary"
        if SUM_SHEET in wb.sheetnames:
            del wb[SUM_SHEET]
        ws_s = wb.create_sheet(SUM_SHEET)
        _write_summary_sheet(ws_s, price_df, financials_df)

        wb.save(path)
        console.print(f"\n  [bold green]✅ Excel saved → {path}[/bold green]")
        console.print(f"  Sheets: {wb.sheetnames}")

    except PermissionError:
        console.print(
            f"  [red]❌ Cannot write to {path}. "
            "Close the file in Excel and retry.[/red]"
        )
    except Exception as e:
        logger.error(f"Excel save error: {e}", exc_info=True)
        console.print(f"  [red]Save error: {e}[/red]")


def _write_summary_sheet(ws, price_df, fin_df):
    """Write a summary overview sheet."""
    from openpyxl.styles import Font, PatternFill, Alignment
    from datetime import datetime

    ws.title = "📊 Summary"
    now      = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    rows = [
        ("Indian Stock Market Scraper – Summary", ""),
        ("Generated At",    now),
        ("", ""),
        ("── Price Data ──", ""),
    ]

    if price_df is not None and not price_df.empty:
        rows += [
            ("Total Price Records", str(len(price_df))),
            ("Indices",     str(len(price_df[price_df["Category"] == "Index"]))),
            ("Top Gainers", str(len(price_df[price_df["Category"] == "Top Gainer"]))),
            ("Top Losers",  str(len(price_df[price_df["Category"] == "Top Loser"]))),
            ("Watchlist",   str(len(price_df[price_df["Category"] == "Watchlist"]))),
            ("Sources",     ", ".join(price_df["Source"].unique())),
        ]
    else:
        rows.append(("Price Data", "Not scraped"))

    rows += [("", ""), ("── Financial Data ──", "")]

    if fin_df is not None and not fin_df.empty:
        rows += [
            ("Companies with Financials", str(len(fin_df))),
            ("Data Source",  "Screener.in"),
            ("Metrics",      "Revenue, PAT, EBITDA (Last 3 Years)"),
            ("Extra Ratios", "PE, ROE, ROCE, D/E, EPS, Book Value"),
        ]
    else:
        rows.append(("Financial Data", "Not scraped"))

    for ri, (label, val) in enumerate(rows, 1):
        c1 = ws.cell(row=ri, column=1, value=label)
        c2 = ws.cell(row=ri, column=2, value=val)
        if ri == 1:
            c1.font = Font(bold=True, size=14, color="1F4E79")
        elif "──" in label:
            c1.font = Font(bold=True, color="2E6DB4")
        c1.alignment = Alignment(horizontal="left")
        c2.alignment = Alignment(horizontal="left")

    ws.column_dimensions["A"].width = 35
    ws.column_dimensions["B"].width = 45


# ─────────────────────────────────────────────────────────────────────────────
# Main application
# ─────────────────────────────────────────────────────────────────────────────

def run_app(config: ScraperConfig, logger):
    dm          = DataManager(config)
    price_df    = None
    fin_df      = None
    filtered_df = None

    while True:
        # Status line
        p_status = f"[green]{len(price_df)} rows[/]"     if price_df    is not None else "[dim]none[/]"
        f_status = f"[green]{len(fin_df)} companies[/]"  if fin_df      is not None else "[dim]none[/]"
        ff_status= f"[green]{len(filtered_df)} rows[/]"  if filtered_df is not None else "[dim]none[/]"

        console.print(Panel(
            f"[bold cyan]📊 Main Menu[/bold cyan]\n"
            f"[dim]Prices: {p_status}  │  "
            f"Financials: {f_status}  │  "
            f"Filtered: {ff_status}[/dim]",
            border_style="blue", padding=(0, 2),
        ))

        items = [
            ("1", "🏢 Select companies & Scrape data"),
            ("2", "📂 Load existing data from Excel"),
            ("3", "🔍 Open Filter Studio"),
            ("4", "📋 View price data table"),
            ("5", "💰 View financial data table"),
            ("6", "📊 View statistics"),
            ("7", "💾 Save all data to Excel"),
            ("8", "🔄 Scrape financials for loaded data"),
            ("9", "⚙  Settings"),
            ("S", "🏭 Sector Analysis Studio (11 sectors, 13 metrics)"),
            ("Q", "🚪 Quit"),
        ]
        for key, label in items:
            colour = "bold green" if key == "Q" else "bold yellow"
            console.print(f"  [{colour}]{key}[/]  {label}")
        console.print()

        choice = Prompt.ask("  Your choice", default="Q").strip().upper()

        # ── 1. Scrape ─────────────────────────────────────────────────────────
        if choice == "1":
            price_df, fin_df = run_scrape_with_selection(config, logger)
            filtered_df = None

        # ── 2. Load from Excel ────────────────────────────────────────────────
        elif choice == "2":
            if not os.path.exists(config.excel_path):
                console.print(f"  [red]File not found: {config.excel_path}[/red]")
                continue
            try:
                wb_sheets = pd.ExcelFile(config.excel_path).sheet_names
                console.print(f"  Available sheets: {wb_sheets}")
                sheet = Prompt.ask("  Sheet to load", default="📋 All Prices")
                loaded = pd.read_excel(config.excel_path, sheet_name=sheet)
                console.print(f"  [green]✔ Loaded {len(loaded)} rows[/green]")

                # Determine if it's price or financial data
                if "Revenue Y1" in loaded.columns:
                    fin_df   = loaded
                    console.print("  → Loaded as Financial Data")
                else:
                    price_df = loaded
                    console.print("  → Loaded as Price Data")

            except Exception as e:
                console.print(f"  [red]Load error: {e}[/red]")

        # ── 3. Filter Studio ──────────────────────────────────────────────────
        elif choice == "3":
            if price_df is None:
                console.print("  [yellow]No price data loaded.[/yellow]")
                continue
            ui          = FilterUI(config)
            filtered_df = ui.run(price_df)
            console.print(
                f"  [green]Filter done: "
                f"{len(price_df)} → {len(filtered_df)} rows[/green]"
            )

        # ── 4. View price table ───────────────────────────────────────────────
        elif choice == "4":
            data = filtered_df if filtered_df is not None else price_df
            if data is None:
                console.print("  [yellow]No price data.[/yellow]")
                continue
            label = "Filtered" if filtered_df is not None else "All"
            max_r = IntPrompt.ask("  Max rows", default=20)
            print_dataframe(data, title=f"Price Data ({label})", max_rows=max_r)

        # ── 5. View financial table ───────────────────────────────────────────
        elif choice == "5":
            if fin_df is None:
                console.print("  [yellow]No financial data.[/yellow]")
                continue
            _show_financials_table(fin_df)

        # ── 6. Statistics ─────────────────────────────────────────────────────
        elif choice == "6":
            data = filtered_df if filtered_df is not None else price_df
            if data is not None:
                print_statistics(data)
            if fin_df is not None:
                _show_financial_stats(fin_df)

        # ── 7. Save to Excel ──────────────────────────────────────────────────
        elif choice == "7":
            if price_df is None and fin_df is None:
                console.print("  [yellow]No data to save.[/yellow]")
                continue
            save_to_excel(config, logger, price_df, fin_df, filtered_df)

        # ── 8. Scrape financials for loaded symbols ───────────────────────────
        elif choice == "8":
            if price_df is None:
                console.print("  [yellow]Load price data first (option 1 or 2).[/yellow]")
                continue
            symbols = [s for s in price_df["Symbol"].unique()
                       if s and s != "N/A" and "^" not in str(s)]
            console.print(
                f"  Will fetch financials for {len(symbols)} symbols: "
                f"{symbols[:5]}…"
            )
            if not Confirm.ask("  Proceed?", default=True):
                continue
            fin_scraper = ScreenerFinancialScraper(config, logger)
            fin_records = []
            with Progress(
                SpinnerColumn(),
                TextColumn("[bold blue]{task.description}"),
                BarColumn(),
                TaskProgressColumn(),
                console=console,
            ) as progress:
                task = progress.add_task("Financials…", total=len(symbols))
                for i, sym in enumerate(symbols, 1):
                    progress.update(
                        task,
                        description=f"[bold blue]{sym} ({i}/{len(symbols)})",
                        completed=i,
                    )
                    fin_records.append(fin_scraper.scrape_company(sym))
            fin_df = pd.DataFrame(fin_records)
            console.print(f"  [green]✔ {len(fin_df)} financial records[/green]")

        # ── 9. Settings ───────────────────────────────────────────────────────
        elif choice == "9":
            _settings_menu(config)
        elif choice == "S":
            sector_ui = SectorUI(config)
            sector_ui.run()


        elif choice == "Q":
            if Confirm.ask("  Exit?", default=True):
                console.print("[bold cyan]\n  Goodbye! Happy Investing! 🇮🇳📈[/bold cyan]\n")
                break
        else:
            console.print(f"  [red]Unknown: '{choice}'[/red]")


def _show_financials_table(fin_df: pd.DataFrame):
    """Display financial data in a rich table."""
    t = Table(
        title       = f"💰 Financial Data ({len(fin_df)} companies)",
        box         = box.ROUNDED,
        border_style= "green",
        header_style= "bold white on dark_green",
        show_lines  = True,
    )
    cols = ["Symbol", "Revenue Y1", "Revenue Y2", "Revenue Y3",
            "PAT Y1", "PAT Y2", "PAT Y3",
            "PE Ratio", "ROE (%)", "ROCE (%)", "Debt to Equity"]
    cols = [c for c in cols if c in fin_df.columns]

    for col in cols:
        t.add_column(col, justify="right" if col != "Symbol" else "left")

    for _, row in fin_df.head(25).iterrows():
        cells = []
        for col in cols:
            val = str(row.get(col, "N/A"))
            if col.startswith("PAT") and val not in ("N/A", ""):
                try:
                    num = float(val)
                    val = f"[{'green' if num > 0 else 'red'}]{val}[/]"
                except ValueError:
                    pass
            cells.append(val)
        t.add_row(*cells)

    console.print(t)


def _show_financial_stats(fin_df: pd.DataFrame):
    """Show summary stats for financial data."""
    console.print(Panel("[bold green]Financial Data Summary[/bold green]",
                        border_style="green"))

    for metric in ["Revenue Y1", "PAT Y1", "PE Ratio", "ROE (%)"]:
        if metric not in fin_df.columns:
            continue
        vals = pd.to_numeric(fin_df[metric], errors="coerce").dropna()
        if vals.empty:
            continue
        console.print(
            f"  [cyan]{metric:<20}[/]  "
            f"Mean: {vals.mean():>12,.2f}  │  "
            f"Max: {vals.max():>12,.2f}  │  "
            f"Min: {vals.min():>12,.2f}"
        )


def _settings_menu(config: ScraperConfig):
    console.print(Panel("[bold]⚙ Settings[/bold]", border_style="dim"))
    console.print(f"  [1] Max stocks       : [cyan]{config.MAX_STOCKS}[/]")
    console.print(f"  [2] Append mode      : [cyan]{config.APPEND_MODE}[/]")
    console.print(f"  [3] Rate limit delay : [cyan]{config.RATE_LIMIT_DELAY}s[/]")
    console.print(f"  [4] Excel path       : [cyan]{config.excel_path}[/]")
    console.print("  [B] Back\n")

    s = Prompt.ask("  Setting", default="B").upper()
    if s == "1":
        config.MAX_STOCKS = IntPrompt.ask("  New max", default=config.MAX_STOCKS)
    elif s == "2":
        config.APPEND_MODE = not config.APPEND_MODE
        console.print(f"  Append mode → {config.APPEND_MODE}")
    elif s == "3":
        val = Prompt.ask("  Delay (seconds)", default=str(config.RATE_LIMIT_DELAY))
        try:
            config.RATE_LIMIT_DELAY = float(val)
        except ValueError:
            pass
    elif s == "4":
        new = Prompt.ask("  New filename", default=config.EXCEL_FILENAME)
        config.EXCEL_FILENAME = new


# ─────────────────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    if sys.version_info < (3, 9):
        print("❌ Python 3.9+ required.")
        sys.exit(1)

    config = ScraperConfig()
    logger = get_logger("main", config)
    show_banner()
    console.print(f"  [dim]Output: {config.excel_path}[/dim]\n")
    run_app(config, logger)