# data_manager.py

import os
from typing import Optional
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import ScraperConfig
from logger_setup import get_logger

HEADER_HEX   = "1F4E79"
GAINER_HEX   = "C6EFCE"
LOSER_HEX    = "FFC7CE"
WATCHLIST_HEX= "FFEB9C"
INDEX_HEX    = "DDEBF7"


class DataManager:
    COLUMNS = [
        "Date","Source","Category","Symbol","Title",
        "Price (₹)","Change","Change (%)","Volume",
        "Market Cap (Cr)","52W High","52W Low","URL",
    ]

    def __init__(self, config: ScraperConfig):
        self.config = config
        self.logger = get_logger(self.__class__.__name__, config)

    def to_dataframe(self, records: list) -> pd.DataFrame:
        if not records:
            return pd.DataFrame(columns=self.COLUMNS)
        df = pd.DataFrame(records)
        for col in self.COLUMNS:
            if col not in df.columns:
                df[col] = "N/A"
        df = df[self.COLUMNS].replace("", "N/A")
        df.drop_duplicates(subset=["Date","Symbol","Title"], inplace=True)
        return df.reset_index(drop=True)

    @staticmethod
    def _style_ws(ws, df):
        thin = Side(style="thin", color="CCCCCC")
        bdr  = Border(left=thin, right=thin, top=thin, bottom=thin)

        # Header
        for ci, col in enumerate(df.columns, 1):
            c = ws.cell(row=1, column=ci, value=col)
            c.font      = Font(bold=True, color="FFFFFF", size=11)
            c.fill      = PatternFill("solid", fgColor=HEADER_HEX)
            c.alignment = Alignment(horizontal="center", wrap_text=True)
            c.border    = bdr

        # Data rows
        cat_colours = {
            "Top Gainer":"C6EFCE", "Top Loser":"FFC7CE",
            "Index":"DDEBF7",      "Watchlist":"FFEB9C",
        }
        for ri, (_, row) in enumerate(df.iterrows(), 2):
            cat  = str(row.get("Category",""))
            fill = PatternFill("solid", fgColor=cat_colours.get(cat,"FFFFFF"))
            for ci, val in enumerate(row, 1):
                c = ws.cell(row=ri, column=ci, value=val)
                c.fill   = fill
                c.border = bdr
                c.alignment = Alignment(wrap_text=False)

        # Column widths
        for col in ws.columns:
            w = max((len(str(c.value or "")) for c in col), default=10)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(w + 3, 45)

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    def save(self, df: pd.DataFrame, filtered_df: Optional[pd.DataFrame] = None):
        if df.empty:
            self.logger.warning("No data to save.")
            return
        path   = self.config.excel_path
        append = self.config.APPEND_MODE and os.path.exists(path)
        try:
            wb = openpyxl.load_workbook(path) if append else openpyxl.Workbook()
            if not append and "Sheet" in wb.sheetnames:
                del wb["Sheet"]

            def write_sheet(name, data):
                if name in wb.sheetnames:
                    ws = wb[name]
                    nr = ws.max_row + 1
                    for ri, (_, row) in enumerate(data.iterrows(), nr):
                        for ci, val in enumerate(row, 1):
                            ws.cell(row=ri, column=ci, value=val)
                else:
                    ws = wb.create_sheet(name)
                    self._style_ws(ws, data)

            write_sheet("📋 All Data", df)

            for cat, sheet_name in [
                ("Index","📈 Indices"),("Top Gainer","🚀 Gainers"),
                ("Top Loser","📉 Losers"),("Watchlist","👁 Watchlist"),
            ]:
                subset = df[df["Category"] == cat]
                if not subset.empty:
                    write_sheet(sheet_name, subset)

            if filtered_df is not None and not filtered_df.empty:
                if "🔍 Filtered" in wb.sheetnames:
                    del wb["🔍 Filtered"]
                ws_f = wb.create_sheet("🔍 Filtered")
                self._style_ws(ws_f, filtered_df)

            wb.save(path)
            self.logger.info(f"✅ Saved → {path}")

        except PermissionError:
            self.logger.error("Close the Excel file and retry.")
        except Exception as e:
            self.logger.error(f"Save error: {e}", exc_info=True)